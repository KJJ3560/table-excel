"""외국인 지분율 상위 종목 + 변동률을 시장별 시트로 나눈 엑셀 리포트를 생성한다."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_COLUMN_LABELS = {
    "rank": "순위",
    "ticker": "종목코드",
    "name": "종목명",
    "foreign_ratio": "지분율(%)",
    "ratio_change": "전일대비(%p)",
    "rank_change": "순위변동",
    "shares_outstanding": "상장주식수",
    "foreign_shares": "외국인보유수량",
    "recent_disclosures": "최근 지분공시(DART)",
}

_RISE_COLOR = "C00000"
_FALL_COLOR = "0000C0"


def _present_columns(df: pd.DataFrame) -> list[str]:
    return [col for col in _COLUMN_LABELS if col in df.columns]


def _style_sheet(ws: Worksheet, columns: list[str]) -> None:
    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    if "전일대비(%p)" in columns:
        change_col_idx = columns.index("전일대비(%p)") + 1
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=change_col_idx)
            if isinstance(cell.value, (int, float)):
                if cell.value > 0:
                    cell.font = Font(color=_RISE_COLOR)
                elif cell.value < 0:
                    cell.font = Font(color=_FALL_COLOR)

    for col_idx, column in enumerate(columns, start=1):
        header_len = len(str(column))
        cell_lengths = [len(str(row[col_idx - 1])) for row in ws.iter_rows(min_row=2, values_only=True)]
        max_len = max([header_len, *cell_lengths]) if cell_lengths else header_len
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)


def build_excel_report(market_frames: dict[str, pd.DataFrame], output_path: Path) -> Path:
    """시장별 DataFrame을 받아 시트로 나눈 엑셀 파일을 생성한다."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for market, df in market_frames.items():
            columns = _present_columns(df)
            sheet_df = df[columns].rename(columns=_COLUMN_LABELS)
            sheet_df.to_excel(writer, sheet_name=market, index=False)
            _style_sheet(writer.sheets[market], list(sheet_df.columns))
    return output_path
